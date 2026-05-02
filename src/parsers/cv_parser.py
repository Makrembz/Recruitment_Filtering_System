import json
import re
import time
from pathlib import Path
from typing import Any

import pdfplumber
from docx import Document
from openai import APIStatusError, RateLimitError


GROQ_MODEL = "llama-3.3-70b-versatile"

CV_EXTRACTION_SYSTEM_PROMPT = """
You are a production-grade multilingual CV parsing engine for a recruitment filtering system.

Your task:
- Extract structured candidate information from CV text.
- The CV may be written in French, English, Arabic, or a mix of these languages.
- Preserve names, institutions, companies, degrees, skills, languages, and certifications exactly when possible.
- Normalize the response into the JSON schema below.
- If a value is missing, use null for scalar fields and [] for list fields.
- duration_months must be an integer or null. Infer it from dates when possible.
- year must be an integer or null.
- raw_text must contain the original extracted CV text exactly as provided.
- Do not invent facts that are not supported by the CV text.

Return ONLY valid JSON with no preamble, markdown, explanations, comments, or trailing text.

Required JSON schema:
{
  "full_name": null,
  "email": null,
  "phone": null,
  "education": [
    {
      "degree": null,
      "institution": null,
      "year": null
    }
  ],
  "experience": [
    {
      "title": null,
      "company": null,
      "duration_months": null,
      "description": null
    }
  ],
  "skills": [],
  "languages": [],
  "certifications": [],
  "raw_text": ""
}
""".strip()

JOB_DESCRIPTION_SYSTEM_PROMPT = """
You are a production-grade multilingual job-description parser for a recruitment filtering system.

Your task:
- Extract hiring requirements from a job description.
- The job description may be written in French, English, Arabic, or a mix of these languages.
- Separate required skills from preferred or nice-to-have skills.
- Infer min_experience_years only when explicitly stated or strongly implied by phrases like "3+ years".
- Keep responsibilities as concise action-oriented strings.
- If a value is missing, use null for scalar fields and [] for list fields.
- Do not invent requirements that are not supported by the text.

Return ONLY valid JSON with no preamble, markdown, explanations, comments, or trailing text.

Required JSON schema:
{
  "required_skills": [],
  "preferred_skills": [],
  "min_experience_years": null,
  "required_education": null,
  "responsibilities": []
}
""".strip()


def parse_cv_file(
    file_path: str | Path,
    client: Any,
    model: str = GROQ_MODEL,
    max_retries: int = 3,
) -> dict[str, Any]:
    """Parse a PDF/DOCX CV and return structured candidate JSON."""
    path = Path(file_path)
    raw_text = extract_text_from_file(path)
    return extract_cv_json(raw_text, client=client, model=model, max_retries=max_retries)


def parse_job_description(
    job_description: str,
    client: Any,
    model: str = GROQ_MODEL,
    max_retries: int = 3,
) -> dict[str, Any]:
    """Parse a multilingual job description into structured hiring criteria."""
    messages = [
        {"role": "system", "content": JOB_DESCRIPTION_SYSTEM_PROMPT},
        {"role": "user", "content": f"JOB DESCRIPTION:\n{job_description}"},
    ]
    try:
        content = _chat_completion_with_retry(
            client=client,
            model=model,
            messages=messages,
            max_retries=max_retries,
            temperature=0,
            max_tokens=900,
        )
        return _normalize_job_payload(_load_json_object(content))
    except Exception:
        return regex_extract_job_description(job_description)


def extract_cv_json(
    raw_text: str,
    client: Any,
    model: str = GROQ_MODEL,
    max_retries: int = 3,
) -> dict[str, Any]:
    """Extract CV JSON with LLM first and regex fallback on failure."""
    messages = [
        {"role": "system", "content": CV_EXTRACTION_SYSTEM_PROMPT},
        {"role": "user", "content": f"CV TEXT:\n{raw_text}"},
    ]
    try:
        content = _chat_completion_with_retry(
            client=client,
            model=model,
            messages=messages,
            max_retries=max_retries,
            temperature=0,
            max_tokens=1800,
        )
        payload = _normalize_cv_payload(_load_json_object(content))
        payload["raw_text"] = raw_text
        return payload
    except Exception:
        return regex_extract_cv(raw_text)


def extract_text_from_file(file_path: str | Path) -> str:
    """Extract raw text from supported CV document formats."""
    path = Path(file_path)
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return extract_text_from_pdf(path)
    if suffix == ".docx":
        return extract_text_from_docx(path)
    raise ValueError(f"Unsupported file type: {suffix}. Expected .pdf or .docx")


def extract_text_from_pdf(file_path: str | Path) -> str:
    chunks: list[str] = []
    with pdfplumber.open(file_path) as pdf:
        for page in pdf.pages:
            text = page.extract_text() or ""
            if text.strip():
                chunks.append(text.strip())
    return "\n\n".join(chunks)


def extract_text_from_docx(file_path: str | Path) -> str:
    document = Document(file_path)
    chunks = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                chunks.append(" | ".join(cells))
    return "\n".join(chunks)


def regex_extract_cv(raw_text: str) -> dict[str, Any]:
    """Best-effort CV fallback when the LLM is unavailable or returns invalid JSON."""
    return _normalize_cv_payload(
        {
            "full_name": _extract_name(raw_text),
            "email": _first_match(EMAIL_RE, raw_text),
            "phone": _first_match(PHONE_RE, raw_text),
            "education": _extract_education(raw_text),
            "experience": _extract_experience(raw_text),
            "skills": _extract_section_items(raw_text, SECTION_LABELS["skills"]),
            "languages": _extract_section_items(raw_text, SECTION_LABELS["languages"]),
            "certifications": _extract_section_items(raw_text, SECTION_LABELS["certifications"]),
            "raw_text": raw_text,
        }
    )


def regex_extract_job_description(job_description: str) -> dict[str, Any]:
    """Best-effort job-description fallback when the LLM is unavailable."""
    skills = _extract_section_items(job_description, SECTION_LABELS["skills"])
    years = _extract_min_years(job_description)
    responsibilities = _extract_section_items(
        job_description, SECTION_LABELS["responsibilities"]
    )
    education = _extract_required_education(job_description)
    return _normalize_job_payload(
        {
            "required_skills": skills,
            "preferred_skills": [],
            "min_experience_years": years,
            "required_education": education,
            "responsibilities": responsibilities,
        }
    )


def _chat_completion_with_retry(
    client: Any,
    model: str,
    messages: list[dict[str, str]],
    max_retries: int,
    temperature: float,
    max_tokens: int,
) -> str:
    for attempt in range(max_retries + 1):
        try:
            response = client.chat.completions.create(
                model=model,
                messages=messages,
                temperature=temperature,
                max_tokens=max_tokens,
                response_format={"type": "json_object"},
            )
            return response.choices[0].message.content or ""
        except (RateLimitError, APIStatusError) as exc:
            if not _is_rate_limit(exc) or attempt >= max_retries:
                raise
            time.sleep(2**attempt)
    raise RuntimeError("LLM request failed after retries")


def _is_rate_limit(exc: Exception) -> bool:
    return isinstance(exc, RateLimitError) or getattr(exc, "status_code", None) == 429


def _load_json_object(content: str) -> dict[str, Any]:
    try:
        payload = json.loads(content)
    except json.JSONDecodeError:
        match = re.search(r"\{.*\}", content, flags=re.DOTALL)
        if not match:
            raise
        payload = json.loads(match.group(0))
    if not isinstance(payload, dict):
        raise ValueError("LLM response must be a JSON object")
    return payload


def _normalize_cv_payload(payload: dict[str, Any]) -> dict[str, Any]:
    return {
        "full_name": _as_optional_string(payload.get("full_name")),
        "email": _as_optional_string(payload.get("email")),
        "phone": _as_optional_string(payload.get("phone")),
        "education": [_normalize_education(item) for item in _as_list(payload.get("education"))],
        "experience": [
            _normalize_experience(item) for item in _as_list(payload.get("experience"))
        ],
        "skills": _string_list(payload.get("skills")),
        "languages": _string_list(payload.get("languages")),
        "certifications": _string_list(payload.get("certifications")),
        "raw_text": str(payload.get("raw_text") or ""),
    }


def _normalize_job_payload(payload: dict[str, Any]) -> dict[str, Any]:
    return {
        "required_skills": _string_list(payload.get("required_skills")),
        "preferred_skills": _string_list(payload.get("preferred_skills")),
        "min_experience_years": _as_optional_int(payload.get("min_experience_years")),
        "required_education": _as_optional_string(payload.get("required_education")),
        "responsibilities": _string_list(payload.get("responsibilities")),
    }


def _normalize_education(item: Any) -> dict[str, Any]:
    item = item if isinstance(item, dict) else {}
    return {
        "degree": _as_optional_string(item.get("degree")),
        "institution": _as_optional_string(item.get("institution")),
        "year": _as_optional_int(item.get("year")),
    }


def _normalize_experience(item: Any) -> dict[str, Any]:
    item = item if isinstance(item, dict) else {}
    return {
        "title": _as_optional_string(item.get("title")),
        "company": _as_optional_string(item.get("company")),
        "duration_months": _as_optional_int(item.get("duration_months")),
        "description": _as_optional_string(item.get("description")),
    }


def _as_optional_string(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _as_optional_int(value: Any) -> int | None:
    if value is None or value == "":
        return None
    try:
        return int(float(str(value).strip()))
    except ValueError:
        return None


def _as_list(value: Any) -> list[Any]:
    if value is None:
        return []
    if isinstance(value, list):
        return value
    return [value]


def _string_list(value: Any) -> list[str]:
    items = []
    for item in _as_list(value):
        if isinstance(item, dict):
            text = item.get("name") or item.get("skill") or item.get("value")
        else:
            text = item
        text = str(text).strip() if text is not None else ""
        if text and text not in items:
            items.append(text)
    return items


EMAIL_RE = re.compile(r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+", re.IGNORECASE)
PHONE_RE = re.compile(r"(?:\+?\d[\d\s().-]{7,}\d)")
YEAR_RE = re.compile(r"\b(?:19|20)\d{2}\b")
MIN_YEARS_RE = re.compile(
    r"(\d+(?:[.,]\d+)?)\s*\+?\s*(?:years?|ans|سن(?:ة|وات)?|عام(?:ا)?|années?)",
    re.IGNORECASE,
)

SECTION_LABELS = {
    "skills": [
        "skills",
        "technical skills",
        "competences",
        "compétences",
        "مهارات",
        "المهارات",
    ],
    "languages": ["languages", "langues", "لغات", "اللغات"],
    "certifications": ["certifications", "certificates", "certificats", "شهادات"],
    "responsibilities": [
        "responsibilities",
        "missions",
        "tasks",
        "responsabilites",
        "responsabilités",
        "مهام",
        "المسؤوليات",
    ],
}

EDUCATION_KEYWORDS = [
    "bachelor",
    "master",
    "phd",
    "doctorate",
    "licence",
    "maitrise",
    "maîtrise",
    "ingenieur",
    "ingénieur",
    "diplome",
    "diplôme",
    "degree",
    "university",
    "universite",
    "université",
    "جامعة",
    "بكالوريوس",
    "ماجستير",
    "دكتوراه",
]


def _first_match(pattern: re.Pattern[str], text: str) -> str | None:
    match = pattern.search(text)
    return match.group(0).strip() if match else None


def _extract_name(raw_text: str) -> str | None:
    for line in raw_text.splitlines()[:10]:
        cleaned = line.strip(" -|\t")
        if not cleaned or EMAIL_RE.search(cleaned) or PHONE_RE.search(cleaned):
            continue
        if 2 <= len(cleaned.split()) <= 5 and len(cleaned) <= 100:
            return cleaned
    return None


def _extract_education(raw_text: str) -> list[dict[str, Any]]:
    results: list[dict[str, Any]] = []
    for line in raw_text.splitlines():
        lower = line.lower()
        if any(keyword in lower for keyword in EDUCATION_KEYWORDS):
            year = _as_optional_int(_first_match(YEAR_RE, line))
            results.append({"degree": line.strip(), "institution": None, "year": year})
    return results[:8]


def _extract_experience(raw_text: str) -> list[dict[str, Any]]:
    results: list[dict[str, Any]] = []
    for line in raw_text.splitlines():
        lower = line.lower()
        if any(token in lower for token in ["experience", "expérience", "work", "emploi", "stage"]):
            results.append(
                {
                    "title": line.strip(),
                    "company": None,
                    "duration_months": None,
                    "description": line.strip(),
                }
            )
    return results[:8]


def _extract_section_items(raw_text: str, labels: list[str]) -> list[str]:
    lines = raw_text.splitlines()
    for index, line in enumerate(lines):
        normalized = line.strip().lower().rstrip(":")
        if any(label.lower() in normalized for label in labels):
            collected = []
            inline = line.split(":", 1)[1] if ":" in line else ""
            if inline:
                collected.extend(_split_items(inline))
            for following in lines[index + 1 : index + 7]:
                if _looks_like_new_section(following):
                    break
                collected.extend(_split_items(following))
            return _dedupe([item for item in collected if item])
    return []


def _split_items(text: str) -> list[str]:
    text = text.strip().strip("-*•")
    parts = re.split(r"[,;|/،؛]|\s+-\s+|•", text)
    return [part.strip(" -*\t") for part in parts if part.strip(" -*\t")]


def _looks_like_new_section(line: str) -> bool:
    stripped = line.strip()
    if not stripped:
        return False
    return len(stripped) <= 40 and stripped.endswith(":")


def _dedupe(items: list[str]) -> list[str]:
    result = []
    for item in items:
        if item not in result:
            result.append(item)
    return result


def _extract_min_years(text: str) -> int | None:
    match = MIN_YEARS_RE.search(text)
    if not match:
        return None
    return _as_optional_int(match.group(1).replace(",", "."))


def _extract_required_education(text: str) -> str | None:
    for line in text.splitlines():
        lower = line.lower()
        if any(keyword in lower for keyword in EDUCATION_KEYWORDS):
            return line.strip()
    return None
